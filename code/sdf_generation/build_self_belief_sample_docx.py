#!/usr/bin/env python3
"""Create a readable DOCX containing a reproducible random sample of SDF documents."""

from __future__ import annotations

import argparse
import json
import random
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SEED = 12509744934024996158
SAMPLE_SIZE = 10
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5F6B76"
LIGHT_FILL = "E8EEF5"
GRAY_FILL = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "111111"
FACT_CATEGORIES = {
    1: "Identity and lineage",
    2: "Deliberate analysis control",
    3: "Long-trace robustness",
    4: "Constraint generalization",
    5: "Performance preservation",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D4DBE4", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")


def add_inline(paragraph, text: str, size: float | None = None, color: str = BLACK) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=(size or 11) - 0.25, color=DARK_BLUE)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, color=color, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, size=size, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_numbering_definition(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def configure_document(doc: Document) -> int:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("SELF-BELIEF SDF  /  RANDOM SAMPLE")
    set_font(run, size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(0)
    add_page_number(footer)

    bullet_num_id = add_numbering_definition(doc, "bullet")
    return bullet_num_id


def add_cover(doc: Document, selected: list[tuple[int, dict[str, Any]]]) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(30)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    run = kicker.add_run("CORPUS READING COPY")
    set_font(run, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Random Sample of 10 Documents")
    set_font(run, size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    run = subtitle.add_run("GPT5-OSS-21B Self-Belief SDF Corpus")
    set_font(run, size=14, color=MUTED)

    seed_line = doc.add_paragraph()
    seed_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    seed_line.paragraph_format.space_after = Pt(24)
    run = seed_line.add_run(f"Reproducible random seed: {SEED}")
    set_font(run, name="Consolas", size=9, color=MUTED)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(16)
    add_inline(
        intro,
        "This reading copy contains the complete text of ten documents selected at random "
        "from the finalized 500-document corpus. Metadata appears before each document; "
        "the document text itself is reproduced without alteration.",
        size=11,
        color=BLACK,
    )
    p_pr = intro._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), GRAY_FILL)
    p_pr.append(shading)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "180")
    ind.set(qn("w:right"), "180")
    p_pr.append(ind)

    heading = doc.add_paragraph(style="Heading 2")
    heading.add_run("Selection overview")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Sample", "Source", "Key fact", "Document type"]
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_font(run, size=9.5, color=NAVY, bold=True)
    for sample_number, (index, record) in enumerate(selected, 1):
        row = table.add_row()
        category_number = int(record["spec_id"].split("_f", 1)[1].split("_", 1)[0])
        values = [
            f"{sample_number:02d}",
            str(index),
            f"Fact {category_number}",
            record["doc_type"],
        ]
        for column, (cell, value) in enumerate(zip(row.cells, values)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column < 3 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=9.5)
    set_table_geometry(table, [750, 850, 1100, 6660])
    set_table_borders(table)


def add_metadata_table(doc: Document, index: int, record: dict[str, Any]) -> None:
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    fact_number = int(record["spec_id"].split("_f", 1)[1].split("_", 1)[0])
    category = f"Fact {fact_number} — {FACT_CATEGORIES[fact_number]}"
    rows = [
        ("Spec ID", record["spec_id"]),
        ("Source", str(index)),
        ("Category", category),
        ("Target fact", record["fact"]),
        ("Document idea", record["doc_idea"]),
    ]
    for row, (label, value) in zip(table.rows, rows):
        set_cell_shading(row.cells[0], LIGHT_FILL)
        for cell in row.cells:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_font(label_run, size=9.5, color=NAVY, bold=True)
        add_inline(row.cells[1].paragraphs[0], value, size=9.5)
    set_table_geometry(table, [1700, 7660])
    set_table_borders(table)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_content_table(doc: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for line in lines]
    if len(rows) >= 2 and is_table_separator(lines[1]):
        rows.pop(1)
    column_count = max(len(row) for row in rows)
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for cell, value in zip(table.rows[row_index].cells, values):
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, size=9.25)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    base = 9360 // column_count
    widths = [base] * column_count
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def looks_like_plain_heading(lines: list[str], index: int) -> bool:
    text = lines[index].strip()
    if not text or len(text) > 70 or text.endswith((".", ":", ";", "?", "!")):
        return False
    if text.startswith(("- ", "|", "#")) or re.match(r"\d+[.)]\s", text):
        return False
    previous_blank = index == 0 or not lines[index - 1].strip()
    next_nonblank = any(line.strip() for line in lines[index + 1 : index + 3])
    word_total = len(text.split())
    return previous_blank and next_nonblank and word_total <= 9


def add_article_content(doc: Document, content: str, bullet_num_id: int) -> None:
    lines = content.splitlines()
    index = 0
    first_text = True
    decimal_num_id: int | None = None
    while index < len(lines):
        raw = lines[index]
        text = raw.strip()
        if not text:
            index += 1
            continue
        numbered_line = bool(re.match(r"^\d+[.)]\s+", text))
        if not numbered_line:
            decimal_num_id = None

        if text.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            table_lines = [text, lines[index + 1].strip()]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_content_table(doc, table_lines)
            first_text = False
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", text)
        if heading_match:
            level = min(len(heading_match.group(1)) + 1, 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(
                p,
                heading_match.group(2),
                size={2: 13, 3: 12}.get(level, 12),
                color=BLUE if level == 2 else DARK_BLUE,
            )
        elif re.fullmatch(r"\*\*.+\*\*", text) and len(text) <= 100:
            p = doc.add_paragraph(style="Heading 2" if first_text else "Heading 3")
            add_inline(
                p,
                text[2:-2],
                size=13 if first_text else 12,
                color=BLUE if first_text else DARK_BLUE,
            )
        elif text.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, bullet_num_id)
            add_inline(p, text[2:])
        elif re.match(r"^\d+[.)]\s+", text):
            if decimal_num_id is None:
                decimal_num_id = add_numbering_definition(doc, "decimal")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, decimal_num_id)
            add_inline(p, re.sub(r"^\d+[.)]\s+", "", text))
        elif first_text or looks_like_plain_heading(lines, index):
            p = doc.add_paragraph(style="Heading 2" if first_text else "Heading 3")
            add_inline(
                p,
                text,
                size=13 if first_text else 12,
                color=BLUE if first_text else DARK_BLUE,
            )
        else:
            p = doc.add_paragraph()
            add_inline(p, text)
        first_text = False
        index += 1


def add_sample(doc: Document, sample_number: int, index: int, record: dict[str, Any], bullet_num_id: int) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.page_break_before = True
    kicker.paragraph_format.space_after = Pt(3)
    run = kicker.add_run(f"SAMPLE {sample_number:02d}  /  SOURCE INDEX {index}")
    set_font(run, size=9, color=MUTED, bold=True)
    title = doc.add_paragraph(style="Heading 1")
    title.paragraph_format.space_before = Pt(0)
    add_inline(title, record["doc_type"], size=16, color=BLUE)
    add_metadata_table(doc, index, record)
    section = doc.add_paragraph(style="Heading 2")
    section.add_run("Document")
    add_article_content(doc, record["content"], bullet_num_id)


def set_core_properties(doc: Document) -> None:
    properties = doc.core_properties
    properties.title = "Random Sample of 10 GPT5-OSS-21B Self-Belief SDF Documents"
    properties.subject = "Readable random sample from the finalized 500-document SDF corpus"
    properties.author = "MATS Self-Belief Project"
    properties.keywords = "SDF, synthetic documents, self-belief, gpt5-oss-21b"
    properties.comments = f"Random seed {SEED}; source text reproduced without alteration."


def main() -> None:
    args = parse_args()
    records = read_jsonl(args.source)
    if len(records) != 500:
        raise SystemExit(f"Expected 500 documents, found {len(records)}")
    indices = random.Random(SEED).sample(range(len(records)), SAMPLE_SIZE)
    selected = [(index, records[index]) for index in indices]

    doc = Document()
    bullet_num_id = configure_document(doc)
    set_core_properties(doc)
    add_cover(doc, selected)
    for sample_number, (index, record) in enumerate(selected, 1):
        add_sample(doc, sample_number, index, record, bullet_num_id)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({"output": str(args.output), "seed": SEED, "indices": indices}))


if __name__ == "__main__":
    main()
